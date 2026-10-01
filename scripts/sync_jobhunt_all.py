import imaplib
import email
from email.header import decode_header
import json
import re
import time
import os
import sys
import datetime
import html
import urllib.request
import gspread
from google.oauth2.service_account import Credentials
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HH_API_TOKEN = "APPLO0LMM7B8JG5P1D6Q3HMEIPTO6HADITJMTP3ENL0O98MFPSL4GTA6IRUQBVF0"
HH_USER_AGENT = "JobApp/1.0 (dmitriyfedotov1908@gmail.com)"
SPREADSHEET_ID = "1rmM-J_XKOuGJAiaB2_McdfP0Wo_5w4NA6Xx3KXXuTq8"
EXCEL_OUTPUT_PATH = r"C:\Users\strel\Desktop\JobHunt_РОП.xlsx"

def clean_html(raw_html):
    if not raw_html:
        return ""
    clean = re.sub(r'<style.*?</style>', '', raw_html, flags=re.DOTALL | re.IGNORECASE)
    clean = re.sub(r'<script.*?</script>', '', clean, flags=re.DOTALL | re.IGNORECASE)
    clean = re.sub(r'<[^>]+>', ' ', clean)
    clean = html.unescape(clean)
    return re.sub(r'\s+', ' ', clean).strip()

def fetch_hh_vacancy(vac_id):
    url = f"https://api.hh.ru/vacancies/{vac_id}"
    headers = {
        "User-Agent": HH_USER_AGENT,
        "Authorization": f"Bearer {HH_API_TOKEN}",
        "Accept": "application/json"
    }
    req = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            name = data.get("name", "Не указана")
            emp = data.get("employer", {}).get("name", "Не указан")
            sal_obj = data.get("salary")
            sal_str = "Не указана"
            if sal_obj:
                f = sal_obj.get("from")
                t = sal_obj.get("to")
                c = sal_obj.get("currency", "RUR")
                sal_str = (f"от {f} " if f else "") + (f"до {t} " if t else "") + c
            alt_url = data.get("alternate_url", f"https://hh.ru/vacancy/{vac_id}")
            return {
                "id": str(vac_id),
                "name": name,
                "employer": emp,
                "salary": sal_str.strip(),
                "url": alt_url
            }
    except Exception as e:
        print(f"Error fetching HH vacancy {vac_id}: {e}")
        return None

def export_to_excel(rows, headers, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "JobHunt_РОП"
    ws.views.sheetView[0].showGridLines = True

    # Header style
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    ws.row_dimensions[1].height = 28

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    status_fills = {
        "🟢 Интервью с ЛПР": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
        "🔵 Сообщение HR": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
        "🔵 Тестовое / Анкета": PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
        "🟡 Откликнулся": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "🔴 Отказ / Неинтересно": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
        "⚪ Новая": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    }

    my_status_idx = headers.index("my_status") if "my_status" in headers else -1

    for row_idx, r in enumerate(rows, start=2):
        ws.append(r)
        ws.row_dimensions[row_idx].height = 22
        st_val = r[my_status_idx] if my_status_idx != -1 and len(r) > my_status_idx else ""
        row_fill = status_fills.get(st_val, None)

        for col_idx in range(1, len(r) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            cell.alignment = align_left
            if col_idx == my_status_idx + 1 and row_fill:
                cell.fill = row_fill
                cell.font = Font(name="Segoe UI", size=10, bold=True)

    # Column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = min(len(val_str), 50)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Freeze header
    ws.freeze_panes = "A2"

    wb.save(output_path)
    print(f"Excel tracker saved to {output_path}")

def sync_all():
    # Working hours check: strictly 09:00 to 18:00 (MSK)
    now = datetime.datetime.now()
    if now.hour < 9 or now.hour > 18 or (now.hour == 18 and now.minute > 0):
        print(f"[{now.strftime('%Y-%m-%d %H:%M:%S')}] Outside working hours (09:00 - 18:00). Skipping sync.")
        return

    print(f"[{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] Starting full email movement sync...")
    
    # 1. Connect to Google Sheets
    sa_path = os.path.join(os.path.dirname(__file__), "service_account.json")
    with open(sa_path, "r", encoding="utf-8") as f:
        sa = json.load(f)

    creds = Credentials.from_service_account_info({
        "type": "service_account",
        "client_email": sa["email"],
        "private_key": sa["privateKey"],
        "token_uri": "https://oauth2.googleapis.com/token"
    }, scopes=["https://www.googleapis.com/auth/spreadsheets"])

    gc = gspread.authorize(creds)
    sh = gc.open_by_key(SPREADSHEET_ID)
    ws_rop = sh.worksheet("JobHunt_РОП")
    ws_rej = sh.worksheet("JobHunt_Rejected")

    rop_rows = ws_rop.get_all_values()
    rej_rows = ws_rej.get_all_values()

    target_headers = [
        "vacancy_id", "created_at", "company", "role", "salary", 
        "url", "status", "pitch_or_reason", "my_status", "last_movement_at", "movement_details"
    ]

    # Ensure header row has all target columns
    if len(rop_rows[0]) < len(target_headers) or rop_rows[0] != target_headers:
        ws_rop.update(range_name="A1", values=[target_headers])
        rop_rows[0] = target_headers
    if len(rej_rows[0]) < len(target_headers) or rej_rows[0] != target_headers:
        ws_rej.update(range_name="A1", values=[target_headers])
        rej_rows[0] = target_headers

    # Pad data rows if needed
    for i in range(1, len(rop_rows)):
        while len(rop_rows[i]) < len(target_headers):
            rop_rows[i].append("")
    for i in range(1, len(rej_rows)):
        while len(rej_rows[i]) < len(target_headers):
            rej_rows[i].append("")

    rop_map = {r[0]: idx for idx, r in enumerate(rop_rows[1:], start=1) if r[0]}
    rej_map = {r[0]: idx for idx, r in enumerate(rej_rows[1:], start=1) if r[0]}

    # 2. Connect to Gmail
    mail = imaplib.IMAP4_SSL("imap.gmail.com")
    mail.login("dmitriyfedotov1908@gmail.com", "ljcztbitmymvyalr")
    mail.select("INBOX")

    status, messages = mail.search(None, "FROM", "hh.ru")
    msg_ids = messages[0].split()

    # Scan recent 150 emails (oldest to newest so latest prevails)
    scan_ids = msg_ids[-150:] if len(msg_ids) > 150 else msg_ids
    print(f"Scanning {len(scan_ids)} recent emails from hh.ru...")

    # Events buffer: vac_id -> latest event dict
    events = {}

    for mid in scan_ids:
        res, data = mail.fetch(mid, "(RFC822)")
        if not data or not data[0]:
            continue
        msg = email.message_from_bytes(data[0][1])

        subj = ""
        for part, enc in decode_header(msg.get("Subject", "")):
            if isinstance(part, bytes):
                subj += part.decode(enc or "utf-8", errors="ignore")
            else:
                subj += str(part)

        lower_subj = subj.lower()
        if any(skip in lower_subj for skip in [
            "вакансии по подписке", "подходящие вакансии", "привлекло внимание", "планируем ближайший год"
        ]):
            continue

        date_tuple = email.utils.parsedate_tz(msg.get("Date"))
        if date_tuple:
            dt = datetime.datetime.fromtimestamp(email.utils.mktime_tz(date_tuple), datetime.timezone.utc)
            # Moscow time UTC+3
            dt_msk = dt + datetime.timedelta(hours=3)
            event_date = dt_msk.strftime("%Y-%m-%d %H:%M")
        else:
            event_date = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

        body = ""
        for p in msg.walk():
            ctype = p.get_content_type()
            if ctype in ["text/plain", "text/html"]:
                payload = p.get_payload(decode=True)
                if payload:
                    cs = p.get_content_charset() or "utf-8"
                    body += payload.decode(cs, errors="ignore")

        clean_text = clean_html(body)
        vac_ids = list(set(re.findall(r"hh\.ru/vacancy/(\d+)", body)))
        
        # Extract direct chat or negotiation link
        chat_links = re.findall(r'href=["\'](https?://hh\.ru/applicant/[^\s"\']+)["\']', body)
        primary_chat = chat_links[0] if chat_links else ""

        # Extract employer message from quotes «...»
        quote_match = re.search(r'«([^»]+)»', clean_text)
        quote_text = quote_match.group(1).strip() if quote_match else ""

        # Event classification
        target_status = None
        details = ""

        if "не готов пригласить" in lower_subj or "отказ" in lower_subj:
            target_status = "🔴 Отказ / Неинтересно"
            details = "Работодатель отклонил отклик"
        elif "приглаш" in lower_subj or "собеседование" in lower_subj or "онлайн-собеседование" in quote_text.lower():
            target_status = "🟢 Интервью с ЛПР"
            details = quote_text if quote_text else "Приглашение на интервью / собеседование"
        elif "сообщение" in lower_subj or "вам написали" in lower_subj or "хотят связаться" in lower_subj:
            if "тестов" in quote_text.lower() or "анкет" in quote_text.lower() or "форма" in quote_text.lower():
                target_status = "🔵 Тестовое / Анкета"
            else:
                target_status = "🔵 Сообщение HR"
            details = quote_text if quote_text else clean_text[:150]
        elif "отклик получен" in lower_subj or "вы откликнулись" in lower_subj or "отклик доставлен" in lower_subj:
            target_status = "🟡 Откликнулся"
            details = "Отклик отправлен работодателю"

        if target_status and vac_ids:
            for vid in vac_ids:
                events[vid] = {
                    "status": target_status,
                    "date": event_date,
                    "details": details[:300] + (f" | Чат: {primary_chat}" if primary_chat else ""),
                    "subject": subj
                }

    mail.logout()
    print(f"Found events for {len(events)} vacancies in recent emails.")

    # 3. Apply events to sheets
    resurrected_from_rej = []
    new_imported = []
    rop_modified = False
    rej_modified = False

    for vid, ev in events.items():
        st = ev["status"]
        dt = ev["date"]
        det = ev["details"]

        # Case 1: In JobHunt_РОП
        if vid in rop_map:
            idx = rop_map[vid]
            row = rop_rows[idx]
            # Update status, date, details
            row[8] = st
            row[9] = dt
            row[10] = det
            rop_modified = True

        # Case 2: In JobHunt_Rejected
        elif vid in rej_map:
            idx = rej_map[vid]
            row = rej_rows[idx]
            # If positive event -> RESURRECT to РОП!
            if st in ["🟢 Интервью с ЛПР", "🔵 Сообщение HR", "🔵 Тестовое / Анкета", "🟡 Откликнулся"]:
                row[8] = st
                row[9] = dt
                row[10] = det
                resurrected_from_rej.append((vid, row))
            else:
                row[8] = st
                row[9] = dt
                row[10] = det
                rej_modified = True

        # Case 3: NOT in either sheet -> Auto-import active vacancies!
        else:
            if st in ["🟢 Интервью с ЛПР", "🔵 Сообщение HR", "🔵 Тестовое / Анкета", "🟡 Откликнулся"]:
                print(f"Importing missing active vacancy {vid} from HH API...")
                hh_info = fetch_hh_vacancy(vid)
                if hh_info:
                    new_row = [
                        hh_info["id"],
                        dt,
                        hh_info["employer"],
                        hh_info["name"],
                        hh_info["salary"],
                        hh_info["url"],
                        "ACTIVE",
                        "Импорт из почты (активный отклик)",
                        st,
                        dt,
                        det
                    ]
                    new_imported.append(new_row)
                    rop_modified = True

    # Perform movements between sheets
    if resurrected_from_rej:
        print(f"Resurrecting {len(resurrected_from_rej)} vacancies from Rejected to JobHunt_РОП!")
        resurrected_ids = set(x[0] for x in resurrected_from_rej)
        # Append to rop
        for vid, row in resurrected_from_rej:
            rop_rows.append(row)
            rop_map[vid] = len(rop_rows) - 1
        # Remove from rej
        rej_rows = [rej_rows[0]] + [r for r in rej_rows[1:] if r[0] not in resurrected_ids]
        rej_map = {r[0]: idx for idx, r in enumerate(rej_rows[1:], start=1) if r[0]}
        rop_modified = True
        rej_modified = True

    if new_imported:
        print(f"Adding {len(new_imported)} newly imported vacancies to JobHunt_РОП!")
        for row in new_imported:
            rop_rows.append(row)
            rop_map[row[0]] = len(rop_rows) - 1
        rop_modified = True

    # 4. Save updates to Google Sheets
    if rop_modified:
        print(f"Updating JobHunt_РОП ({len(rop_rows)} rows)...")
        ws_rop.clear()
        ws_rop.update(range_name="A1", values=rop_rows)
        print("JobHunt_РОП updated successfully.")

    if rej_modified:
        print(f"Updating JobHunt_Rejected ({len(rej_rows)} rows)...")
        ws_rej.clear()
        ws_rej.update(range_name="A1", values=rej_rows)
        print("JobHunt_Rejected updated successfully.")

    # 5. Export to local Excel
    print("Generating local Excel tracker JobHunt_РОП.xlsx...")
    export_to_excel(rop_rows[1:], target_headers, EXCEL_OUTPUT_PATH)
    print(f"[{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] Full sync completed successfully!")

if __name__ == "__main__":
    sync_all()
